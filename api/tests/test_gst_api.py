"""GST reconciliation over HTTP: permissions, isolation, and the whole flow."""

from __future__ import annotations

import io
import json

import pytest

from api.tests.conftest import member, sign_in
from core.identifiers import gstin_check_character
from core.models import Role
from core.provisioning import create_firm

pytestmark = [pytest.mark.django_db, pytest.mark.usefixtures("fixture_adapters")]

V1 = "/api/v1"


def _g(first):
    return first + gstin_check_character(first)


ME = _g("27AAAPL1234C1Z")
ACME = _g("29ABCPE1234F1Z")

REGISTER = (
    "GSTIN,Invoice No,Date,Taxable Value,CGST,SGST\n"
    f"{ACME},INV-1,12-08-2026,10000,900,900\n"
    f"{ACME},INV-3,25-08-2026,2000,180,180\n"
).encode()

PORTAL = json.dumps(
    {"data": {"gstin": ME, "rtnprd": "082026", "docdata": {"b2b": [{"ctin": ACME, "inv": [
        {"inum": "INV-1", "dt": "12-08-2026", "txval": 10000, "cgst": 900, "sgst": 900}]}]}}}
).encode()


def _file(name, data):
    f = io.BytesIO(data)
    f.name = name
    return {"file": f}


def _start(api, client_record):
    base = f"{V1}/clients/{client_record.pk}/gst"
    reg = api.post(f"{base}/registrations/", {"gstin": ME}, format="json")
    assert reg.status_code == 201, reg.content
    run = api.post(f"{base}/runs/", {"registration": reg.json()["id"], "period": "2026-08"}, format="json")
    assert run.status_code == 201, run.content
    return base, run.json()["id"]


def test_the_whole_flow_over_http(api, client_record):
    base, run = _start(api, client_record)
    assert api.post(f"{base}/runs/{run}/register/", _file("r.csv", REGISTER), format="multipart").status_code == 200
    assert api.post(f"{base}/runs/{run}/portal/", _file("2b.json", PORTAL), format="multipart").status_code == 200
    body = api.post(f"{base}/runs/{run}/reconcile/").json()

    assert body["summary"]["eligible_paise"] == 1800_00
    assert [g["kind"] for g in body["groups"]] == ["missing_in_2b", "matched"]
    assert any("Follow up" in a["text"] for a in body["actions"])

    out = api.get(f"{base}/runs/{run}/export/")
    assert out.status_code == 200 and out.content[:2] == b"PK"  # an xlsx is a zip

    signed = api.post(f"{base}/runs/{run}/sign-off/")
    assert signed.status_code == 200 and signed.json()["status"] == "signed_off"
    frozen = api.post(f"{base}/runs/{run}/reconcile/")
    assert frozen.status_code == 409 and frozen.json()["code"] == "gst_rule"


def test_a_file_that_is_not_2b_is_a_422_with_a_reason(api, client_record):
    base, run = _start(api, client_record)
    r = api.post(f"{base}/runs/{run}/portal/", _file("2b.json", b'{"a": 1}'), format="multipart")
    assert r.status_code == 422 and r.json()["code"] == "gst_file_unreadable"


def test_staff_can_prepare_but_not_sign_off(firm, client_record, staff_api, api):
    base, run = _start(staff_api, client_record)
    staff_api.post(f"{base}/runs/{run}/register/", _file("r.csv", REGISTER), format="multipart")
    staff_api.post(f"{base}/runs/{run}/portal/", _file("2b.json", PORTAL), format="multipart")
    staff_api.post(f"{base}/runs/{run}/reconcile/")
    assert staff_api.post(f"{base}/runs/{run}/sign-off/").status_code == 403


def test_a_reader_can_look_but_not_change(firm, client_record, api):
    _start(api, client_record)
    reader = sign_in(member(firm, Role.READ_ONLY, email="viewer@example.com").user)
    base = f"{V1}/clients/{client_record.pk}/gst"
    assert reader.get(f"{base}/registrations/").status_code == 200
    assert reader.post(f"{base}/registrations/", {"gstin": ACME}, format="json").status_code == 403


def test_another_firms_client_and_runs_are_not_found(api, client_record):
    other = create_firm("Other Firm")
    intruder = sign_in(member(other, Role.SENIOR_CA, email="x@other.example").user)
    base, run = _start(api, client_record)
    assert intruder.get(f"{base}/runs/{run}/").status_code == 404
    assert intruder.get(f"{base}/registrations/").status_code == 404


def test_the_working_paper_never_holds_a_live_formula_or_an_illegal_character(api, client_record):
    from openpyxl import load_workbook

    base, run = _start(api, client_record)
    register = (
        "GSTIN,Invoice No,Date,Taxable Value,CGST,SGST,Supplier Name\n"
        f'{ACME},=1+1,12-08-2026,10000,900,900,"=HYPERLINK(""http://x.example"",""go"")"\n'
        f"{ACME},+SUM(A1),13-08-2026,100,9,9,-2+3\n"
        f"{ACME},@cmd,14-08-2026,100,9,9,\tTabbed\n"
        f"{ACME},INV-9,15-08-2026,100,9,9,Bad\x01Name\n"
    ).encode()
    assert api.post(f"{base}/runs/{run}/register/", _file("r.csv", register), format="multipart").status_code == 200
    assert api.post(f"{base}/runs/{run}/portal/", _file("2b.json", PORTAL), format="multipart").status_code == 200
    assert api.post(f"{base}/runs/{run}/reconcile/").status_code == 200

    out = api.get(f"{base}/runs/{run}/export/")

    assert out.status_code == 200
    wb = load_workbook(io.BytesIO(out.content))
    cells = [c for ws in wb for row in ws.iter_rows() for c in row if c.value is not None]
    assert not [c for c in cells if c.data_type == "f"]
    text = [str(c.value) for c in cells]
    assert "'=1+1" in text and "'@cmd" in text
    assert any(t == "BadName" for t in text), [t for t in text if "Bad" in t or "INV" in t or "Tab" in t]


@pytest.mark.parametrize("mapping", ["[1]", '"x"', "5", '{"invoice_no": 3}'])
def test_a_column_mapping_that_is_not_an_object_of_strings_is_a_422(api, client_record, mapping):
    base, run = _start(api, client_record)
    r = api.post(
        f"{base}/runs/{run}/register/", {**_file("r.csv", REGISTER), "mapping": mapping}, format="multipart"
    )
    assert r.status_code == 422 and r.json()["code"] == "gst_file_unreadable"


@pytest.mark.parametrize(
    "document",
    [
        b'{"data":{"docdata":{"b2b":5}}}',
        b'{"data":{"docdata":{"b2b":["x"]}}}',
        b'{"data":{"docdata":{"b2b":[{"inv":"x"}]}}}',
        b'{"data":{"docdata":{"b2b":[{"inv":[{"items":3}]}]}}}',
        b'{"data":{"docdata":{"b2b":[{"ctin":[1],"inv":[{"txval":[1]}]}]}}}',
        b'{"data":{"docdata":[]}}',
        b"[" * 200_000,
    ],
    ids=["b2b-number", "b2b-string", "supplier-not-doc", "items-number", "odd-types", "docdata-list", "deep-nesting"],
)
def test_malformed_gstr2b_shapes_are_a_422_not_a_500(api, client_record, document):
    base, run = _start(api, client_record)
    r = api.post(f"{base}/runs/{run}/portal/", _file("2b.json", document), format="multipart")
    assert r.status_code == 422 and r.json()["code"] == "gst_file_unreadable"


def test_the_registrations_list_is_the_paginated_envelope(api, client_record):
    _start(api, client_record)
    body = api.get(f"{V1}/clients/{client_record.pk}/gst/registrations/").json()

    assert set(body) == {"count", "next", "previous", "results"}
    assert body["count"] == 1 and body["results"][0]["gstin"] == ME


def test_the_schema_documents_the_gst_and_assistant_responses():
    from drf_spectacular.generators import SchemaGenerator

    paths = SchemaGenerator().get_schema(request=None, public=True)["paths"]
    base = "/api/v1/clients/{client_id}/gst/"
    wanted = [
        (f"{base}registrations/", "get"), (f"{base}registrations/", "post"),
        (f"{base}runs/", "get"), (f"{base}runs/", "post"), (f"{base}runs/{{id}}/", "get"),
        (f"{base}runs/{{id}}/register/", "post"), (f"{base}runs/{{id}}/portal/", "post"),
        (f"{base}runs/{{id}}/reconcile/", "post"), (f"{base}runs/{{id}}/decisions/", "post"),
        (f"{base}runs/{{id}}/export/", "get"), (f"{base}runs/{{id}}/sign-off/", "post"),
        ("/api/v1/clients/{client_id}/assistant/next-batch/", "post"),
    ]
    for path, method in wanted:
        responses = paths[path][method]["responses"]
        ok = next(v for code, v in responses.items() if code.startswith("2"))
        assert ok["content"], (path, method)
    run_body = paths[f"{base}runs/"]["post"]["requestBody"]["content"]["application/json"]["schema"]["$ref"]
    assert run_body.endswith("RunCreateRequest")
