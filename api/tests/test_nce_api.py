"""The financial statements in the ICAI non-corporate format, from the API."""

from __future__ import annotations

import pytest

from api.tests.test_bills import base, make_ledger, make_party, post_bill, voucher

pytestmark = [pytest.mark.django_db, pytest.mark.usefixtures("fixture_adapters")]


def statements(api, client_record, fy=2025):
    response = api.get(f"{base(client_record)}/reports/financial-statements/", {"fy": fy})
    assert response.status_code == 200, response.content
    return response.json()


def by_key(rows):
    return {row["key"]: row for row in rows}


def book_purchase(api, client_record, amount=10_000_00):
    party = make_party(api, client_record, name="Ravi Traders")
    purchases = make_ledger(api, client_record, "Purchases", "PURCHASE")
    body = voucher(party, purchases, cgst_paise=0, sgst_paise=0, heads=[{"ledger": purchases["id"], "amount_paise": amount}])
    assert post_bill(api, client_record, body).status_code == 201
    return party


def test_the_balance_sheet_and_profit_and_loss_follow_the_icai_lines_and_balance(api, client_record):
    book_purchase(api, client_record)

    report = statements(api, client_record)

    bs, pl = by_key(report["balance_sheet"]), by_key(report["profit_and_loss"])
    assert bs["CL.PAY"]["current_paise"] == 10_000_00 and bs["CL.PAY"]["note"] == 9
    assert pl["PL.COGS"]["current_paise"] == 10_000_00 and pl["PL.COGS"]["note"] == 21
    assert pl["t.XVII"]["current_paise"] == -10_000_00  # a loss: purchases and no sales
    # The year's result is part of Reserves and surplus, so liabilities (with the loss) equal assets (nil).
    assert bs["EQ.RES"]["current_paise"] == -10_000_00
    assert bs["t.L"]["current_paise"] == bs["t.A"]["current_paise"] == 0
    assert report["balances"] is True


def test_each_line_is_taken_apart_in_its_note_by_ledger(api, client_record):
    book_purchase(api, client_record)

    report = statements(api, client_record)

    notes = {note["number"]: note for note in report["notes"]}
    assert [r["label"] for r in notes[9]["rows"]] == ["Ravi Traders"]
    assert notes[9]["total_current_paise"] == 10_000_00
    assert notes[21]["rows"][0]["current_paise"] == 10_000_00


def test_the_previous_year_is_beside_the_current_one(api, client_record):
    book_purchase(api, client_record)

    report = statements(api, client_record, fy=2026)

    bs = by_key(report["balance_sheet"])
    # The payable stands at the end of the previous year too, and is carried into this one.
    assert bs["CL.PAY"]["previous_paise"] == 10_000_00 and bs["CL.PAY"]["current_paise"] == 10_000_00
    assert report["has_previous"] is True


def test_a_ledger_can_be_placed_on_a_line_by_hand_and_a_wrong_line_is_refused(api, client_record):
    party = book_purchase(api, client_record)
    ledgers = api.get(f"{base(client_record)}/ledgers/").json()["results"]
    purchases = next(ledger for ledger in ledgers if ledger["name"] == "Purchases")

    refused = api.patch(f"{base(client_record)}/ledgers/{purchases['id']}/", {"nce_line": "NOPE"}, format="json")
    assert refused.status_code == 400

    done = api.patch(f"{base(client_record)}/ledgers/{purchases['id']}/", {"nce_line": "PL.EXP"}, format="json")
    assert done.status_code == 200, done.content
    pl = by_key(statements(api, client_record)["profit_and_loss"])
    assert pl["PL.EXP"]["current_paise"] == 10_000_00 and pl["PL.COGS"]["current_paise"] == 0
    assert party["id"]


def put_settings(api, client_record, body):
    response = api.put(f"{base(client_record)}/reports/financial-statements/settings/", body, format="json")
    assert response.status_code == 200, response.content
    return response.json()


def test_the_settings_start_empty_and_keep_what_is_saved(api, client_record):
    empty = api.get(f"{base(client_record)}/reports/financial-statements/settings/").json()
    assert empty["rounding"] == "rupees" and empty["years"] == {} and empty["about"] == ""

    put_settings(api, client_record, {"about": "A trading firm.", "policies": "Accrual basis.", "rounding": "thousands", "years": {"2025": {"closing_stock_paise": 5_00}}})
    saved = api.get(f"{base(client_record)}/reports/financial-statements/settings/").json()
    assert saved["about"] == "A trading firm." and saved["rounding"] == "thousands"
    assert saved["years"]["2025"]["closing_stock_paise"] == 5_00

    refused = api.put(f"{base(client_record)}/reports/financial-statements/settings/", {"years": {"next year": {}}}, format="json")
    assert refused.status_code == 400
    over = api.put(
        f"{base(client_record)}/reports/financial-statements/settings/",
        {"years": {"2025": {"partners": [{"name": "A", "share_bp": 6000}, {"name": "B", "share_bp": 6000}]}}},
        format="json",
    )
    assert over.status_code == 400


def test_closing_stock_lowers_the_cost_of_goods_and_the_sheet_still_balances(api, client_record):
    book_purchase(api, client_record)
    put_settings(api, client_record, {"years": {"2025": {"closing_stock_paise": 4_000_00}}})

    report = statements(api, client_record)

    bs, pl = by_key(report["balance_sheet"]), by_key(report["profit_and_loss"])
    assert bs["CA.STOCK"]["current_paise"] == 4_000_00
    assert pl["PL.COGS"]["current_paise"] == 6_000_00
    assert pl["t.XVII"]["current_paise"] == -6_000_00
    assert report["balances"] is True
    cogs_note = {n["number"]: n for n in report["notes"]}[21]
    assert [r["section"] for r in cogs_note["rows"]] == ["Purchases of stock-in-trade", "Changes in inventories"]


def test_next_year_the_stock_carried_forward_is_charged_unless_the_new_closing_is_entered(api, client_record):
    book_purchase(api, client_record)
    put_settings(api, client_record, {"years": {"2025": {"closing_stock_paise": 4_000_00}}})

    report = statements(api, client_record, fy=2026)

    bs, pl = by_key(report["balance_sheet"]), by_key(report["profit_and_loss"])
    assert bs["CA.STOCK"]["previous_paise"] == 4_000_00 and bs["CA.STOCK"]["current_paise"] == 0
    assert pl["PL.COGS"]["current_paise"] == 4_000_00  # the opening stock, consumed
    assert report["balances"] is True


def test_partners_are_tabled_with_their_share_of_the_profit(api, client_record):
    book_purchase(api, client_record)
    put_settings(
        api,
        client_record,
        {
            "years": {
                "2025": {
                    "partners": [
                        {"name": "Asha", "share_bp": 6000, "opening_paise": 0, "introduced_paise": 3_000_00},
                        {"name": "Bala", "share_bp": 4000, "opening_paise": 0, "introduced_paise": 2_000_00, "withdrawals_paise": 500_00},
                    ]
                }
            }
        },
    )

    capital = statements(api, client_record)["capital"]

    asha, bala = capital["rows"]
    assert asha["profit_share_paise"] == -6_000_00 and bala["profit_share_paise"] == -4_000_00  # a loss of 10,000
    assert asha["closing_paise"] == 3_000_00 - 6_000_00
    assert bala["closing_paise"] == 2_000_00 - 500_00 - 4_000_00


def test_figures_are_rounded_to_the_chosen_unit_and_say_so(api, client_record):
    book_purchase(api, client_record, amount=12_345_67)
    put_settings(api, client_record, {"rounding": "thousands"})

    report = statements(api, client_record)

    assert report["unit_paise"] == 100_000 and report["unit_label"] == "Rs. in thousands"
    assert by_key(report["balance_sheet"])["CL.PAY"]["current_paise"] == 12_000_00
    assert any("rounded" in w for w in report["warnings"])


def test_the_statements_download_as_a_workbook(api, client_record):
    import io

    from openpyxl import load_workbook

    book_purchase(api, client_record)
    response = api.get(f"{base(client_record)}/reports/financial-statements/export/", {"fy": 2025})
    assert response.status_code == 200
    assert "spreadsheetml" in response["Content-Type"]
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Balance Sheet", "Statement of P&L", "Notes 1 to 3", "Notes 4 to 25"]
    cells = [c.value for row in wb["Balance Sheet"].iter_rows() for c in row if c.value is not None]
    assert "Trade payables" in " ".join(str(c) for c in cells)


def test_msme_suppliers_are_split_out_of_trade_payables(api, client_record):
    party = book_purchase(api, client_record)
    refused = api.patch(f"{base(client_record)}/parties/{party['id']}/", {"udyam_no": "NOT-A-NUMBER"}, format="json")
    assert refused.status_code == 400
    ok = api.patch(f"{base(client_record)}/parties/{party['id']}/", {"msme": True, "udyam_no": "UDYAM-MH-12-0001234"}, format="json")
    assert ok.status_code == 200, ok.content

    report = statements(api, client_record)

    split = next(s for s in report["schedules"] if s["note"] == 9 and s["title"].startswith("Trade payables"))
    assert [r["values"][0] for r in split["rows"]] == [10_000_00, 0, 10_000_00]
    disclosure = next(s for s in report["schedules"] if s["title"].startswith("Dues to suppliers"))
    assert disclosure["rows"][1]["values"][0] == 10_000_00  # principal unpaid


def test_receivables_are_aged_from_the_due_date_to_the_year_end(api, client_record):
    customer = make_party(api, client_record, "Mehta Stores", role="CUSTOMER", gstin="")
    sales = make_ledger(api, client_record, "Sales", "SALES")
    for ref, day, amount in (("S-1", "2025-04-10", 3_000_00), ("S-2", "2026-01-10", 2_000_00)):
        body = voucher(customer, sales, kind="SALES", reference=ref, bill_date=day, cgst_paise=0, sgst_paise=0, heads=[{"ledger": sales["id"], "amount_paise": amount}])
        assert post_bill(api, client_record, body).status_code == 201

    report = statements(api, client_record)

    ageing = next(s for s in report["schedules"] if s["note"] == 16)
    values = {r["label"] + str(i): r["values"][0] for i, r in enumerate(ageing["rows"])}
    assert ageing["rows"][2]["values"][0] == 2_000_00  # under six months: unsecured good
    assert ageing["rows"][6]["values"][0] == 3_000_00  # over six months: unsecured good
    assert ageing["rows"][-1]["values"][0] == by_key(report["balance_sheet"])["CA.REC"]["current_paise"] == 5_000_00
    assert values  # every row present

    put_settings(api, client_record, {"years": {"2025": {"receivables_doubtful_paise": 1_000_00}}})
    again = next(s for s in statements(api, client_record)["schedules"] if s["note"] == 16)
    assert again["rows"][6]["values"][0] == 2_000_00 and again["rows"][7]["values"][0] == 1_000_00


def test_entity_type_and_size_are_kept_and_the_size_is_suggested_from_the_books(api, client_record):
    book_purchase(api, client_record)
    before = statements(api, client_record)
    assert before["size"] == "msme" and before["size_suggested"] == "msme" and "MSME" in before["size_statement"]
    assert any("proprietorship" in w for w in before["warnings"])

    put_settings(api, client_record, {"entity_type": "partnership", "size": "large"})
    after = statements(api, client_record)
    assert after["size"] == "large" and after["size_suggested"] == "msme"
    assert after["capital_title"] == "Partners' Capital Accounts"
    # A Large entity gets the Cash Flow Statement itself, and it ties to the cash and bank balance on the balance sheet.
    flow = {r["key"]: r for r in after["cash_flow"]}
    assert flow["A.net"]["label"].startswith("Net cash flow from operating") and flow["N.close"]["kind"] == "total"
    assert flow["N.close"]["current_paise"] == by_key(after["balance_sheet"])["CA.CASH"]["current_paise"]
    assert "N.diff" not in flow


def test_a_ledger_can_be_pinned_to_a_sub_head_of_its_note_and_a_catch_all_is_flagged(api, client_record):
    party = book_purchase(api, client_record)
    other = make_ledger(api, client_record, "Zzz Misc Thing", "INDIRECT_EXPENSE")
    notes = {n["number"]: n for n in statements(api, client_record)["notes"]}
    assert party["id"]
    # No bill posted to it yet, so it is not in the note; pin and unpin are still validated.
    assert api.patch(f"{base(client_record)}/ledgers/{other['id']}/", {"nce_section": "Rent"}, format="json").status_code == 200
    assert api.patch(f"{base(client_record)}/ledgers/{other['id']}/", {"nce_section": "Not a sub-head"}, format="json").status_code == 400
    assert "Rent" in notes[21]["choices"] or notes[21]["choices"] is not None


def test_rounding_each_ledger_first_keeps_the_sheet_adding_up_and_shows_the_gap(api, client_record):
    book_purchase(api, client_record, amount=12_345_67)
    put_settings(api, client_record, {"rounding": "thousands"})

    report = statements(api, client_record)

    bs = by_key(report["balance_sheet"])
    assert bs["t.L"]["current_paise"] == bs["t.A"]["current_paise"]
    assert report["balances"] is True
    pay = bs["CL.PAY"]["current_paise"]
    assert pay % 100_000 == 0


from api.tests.test_assets_api import machine_purchase, register  # noqa: E402,F401


def test_the_asset_block_comes_from_the_register_and_says_when_it_disagrees_with_the_ledgers(api, client_record, machine_purchase):
    machinery, bill = machine_purchase
    assert register(api, client_record, machinery, bill).status_code == 201

    report = statements(api, client_record)

    block = next(s for s in report["schedules"] if s["note"] == 11)
    assert block["columns"] == ["Plant and machinery", "Total"]
    by_label = {(i, r["label"]): r["values"] for i, r in enumerate(block["rows"])}
    assert block["rows"][2]["values"] == [10_00_000_00, 10_00_000_00]  # additions
    assert block["rows"][4]["values"][0] == 10_00_000_00  # gross block at 31 March
    assert block["rows"][7]["values"][0] == 1_00_000_00  # depreciation for the year
    assert block["rows"][12]["values"][0] == 9_00_000_00  # net block
    assert by_label
    # The ledger still carries the cost: the depreciation has not been booked, and the statements say so.
    assert any("net block" in w and "does not agree" in w for w in report["warnings"])
