"""Reading the two inputs: GSTR-2B from the portal, and the purchase register.

Both become lists of ``gst.matching.Invoice`` and nothing else -- the matcher
never learns which file format an invoice came from.

**Refusing is better than guessing.** A row whose amount cannot be read, or a
file with no recognisable invoice-number column, raises ``GstParseError`` with
the row and the reason. A silently skipped row is a missing invoice in the
reconciliation, which reads as "supplier hasn't filed" -- the one wrong answer
here that looks like a perfectly ordinary right one.

Money crosses this boundary as ``Decimal`` parsed from the file's own text (JSON
is read with ``parse_float=Decimal``), never as a float, then to paise through
``core.money``. Spreadsheet cells are the one place a float is unavoidable; a
float that is a whole number of paise is accepted, and any residue finer than a
paisa is rounded away here, explicitly, as the float artefact it is.
"""

from __future__ import annotations

import csv
import datetime
import io
import json
import re
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from core.money import MoneyError, to_paise
from gst.matching import Invoice, Section, normalise_gstin


class GstParseError(ValueError):
    """The file cannot be read as the thing it claims to be."""


@dataclass(frozen=True)
class ParsedPortal:
    gstin: str  # the recipient: whose 2B this is
    period_start: datetime.date | None
    invoices: list[Invoice]


# ---------------------------------------------------------------------------
# Values
# ---------------------------------------------------------------------------


def _paise(value, where: str) -> int:
    if value is None or value == "":
        return 0
    try:
        if isinstance(value, float):
            value = Decimal(repr(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        return to_paise(value)
    except (MoneyError, InvalidOperation, TypeError, ValueError) as exc:
        raise GstParseError(f"{where}: {exc}") from exc


_DATE_FORMATS = ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d/%m/%y", "%d-%b-%y")


def _date(value, where: str) -> datetime.date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise GstParseError(f"{where}: cannot read {text!r} as a date")


def _truthy(value) -> bool:
    return str(value or "").strip().lower() in {"y", "yes", "true", "1", "rcm"}


# ---------------------------------------------------------------------------
# GSTR-2B (JSON, as the portal provides it)
# ---------------------------------------------------------------------------


def parse_gstr2b_json(data: bytes) -> ParsedPortal:
    try:
        doc = json.loads(data.decode("utf-8-sig"), parse_float=Decimal)
    except (UnicodeDecodeError, json.JSONDecodeError, RecursionError) as exc:
        raise GstParseError(f"Not a JSON file: {exc.__class__.__name__}") from exc
    body = doc.get("data") if isinstance(doc, dict) else None
    if not isinstance(body, dict) or "docdata" not in body:
        raise GstParseError("This does not look like a GSTR-2B download (no data.docdata).")
    if not isinstance(body["docdata"], dict):
        raise GstParseError("GSTR-2B: data.docdata should be an object of sections.")

    period = body.get("rtnprd")
    period_start = None
    if isinstance(period, str) and re.fullmatch(r"\d{6}", period) and 1 <= int(period[:2]) <= 12:
        period_start = datetime.date(int(period[2:]), int(period[:2]), 1)

    invoices: list[Invoice] = []
    docdata = _Sections(body["docdata"])
    # b2b and b2ba hold invoices; cdnr and cdnra hold credit (C) and debit (D)
    # notes, told apart by "typ". The "a" sections are amendments, and carry the
    # number they replace as oinum / ontnum.
    for section_key, list_key, number_key, original_key in (
        ("b2b", "inv", "inum", ""),
        ("b2ba", "inv", "inum", "oinum"),
        ("cdnr", "nt", "ntnum", ""),
        ("cdnra", "nt", "ntnum", "ontnum"),
    ):
        for supplier in docdata.objects(section_key):
            for doc_ in _objects(supplier.get(list_key), f"{section_key}.{list_key}"):
                is_note = list_key == "nt"
                section = Section.B2B
                if is_note:
                    section = Section.DN if str(doc_.get("typ", "C")).upper() == "D" else Section.CDN
                invoices.append(
                    _portal_invoice(
                        supplier, doc_, number_key, section, len(invoices) + 1, original_key
                    )
                )
    for supplier in docdata.objects("isd"):
        for doc_ in _objects(supplier.get("doclist") or supplier.get("inv"), "isd.doclist"):
            invoices.append(
                _portal_invoice(
                    supplier, doc_, "docnum", Section.ISD, len(invoices) + 1, "", date_key="docdt"
                )
            )
    for n, doc_ in enumerate(docdata.objects("impg"), start=len(invoices) + 1):
        invoices.append(
            Invoice(
                ref=f"2b:{n}",
                gstin="",
                invoice_no=str(doc_.get("boenum", doc_.get("beno", ""))).strip(),
                invoice_date=_date(doc_.get("boedt", doc_.get("bedt")), f"GSTR-2B entry {n}"),
                taxable_paise=_paise(doc_.get("txval"), f"GSTR-2B entry {n} txval"),
                igst_paise=_paise(doc_.get("igst"), f"GSTR-2B entry {n} igst"),
                cess_paise=_paise(doc_.get("cess"), f"GSTR-2B entry {n} cess"),
                supplier_name=f"Port {doc_.get('portcode', doc_.get('portcd', ''))}".strip(),
                section=Section.IMPG,
                hsn="",
            )
        )

    return ParsedPortal(
        gstin=normalise_gstin(str(body.get("gstin") or "")), period_start=period_start, invoices=invoices
    )


def _objects(value, where: str) -> list[dict]:
    """``value`` as a list of JSON objects, or a plain refusal. Absent is empty."""
    if value is None or value == []:
        return []
    if not isinstance(value, list) or not all(isinstance(v, dict) for v in value):
        raise GstParseError(f"GSTR-2B: {where} should be a list of objects.")
    return value


class _Sections:
    def __init__(self, docdata: dict):
        self._docdata = docdata

    def objects(self, key: str) -> list[dict]:
        return _objects(self._docdata.get(key), key)


def _portal_invoice(
    supplier: dict,
    doc: dict,
    number_key: str,
    section: Section,
    n: int,
    original_key: str = "",
    *,
    date_key: str = "dt",
) -> Invoice:
    where = f"GSTR-2B entry {n}"
    # The portal reports amounts either on the document or itemised; support both.
    src = _objects(doc.get("items"), f"{where} items") or [doc]
    total = dict.fromkeys(("txval", "igst", "cgst", "sgst", "cess"), 0)
    for item in src:
        for k in total:
            total[k] += _paise(item.get(k), f"{where} {k}")
    return Invoice(
        ref=f"2b:{n}",
        gstin=normalise_gstin(str(supplier.get("ctin") or "")),
        invoice_no=str(doc.get(number_key, "")).strip(),
        invoice_date=_date(doc.get(date_key), where),
        taxable_paise=total["txval"],
        igst_paise=total["igst"],
        cgst_paise=total["cgst"],
        sgst_paise=total["sgst"],
        cess_paise=total["cess"],
        supplier_name=str(supplier.get("trdnm", "")).strip(),
        section=section,
        rcm=_truthy(doc.get("rev")),
        # The portal's own verdict on whether the credit can be taken.
        itc_available=str(doc.get("itcavl", "Y")).strip().upper() != "N",
        itc_reason=str(doc.get("rsn", "") or "")[:16],
        original_no=str(doc.get(original_key, "") or "").strip() if original_key else "",
    )


# ---------------------------------------------------------------------------
# Spreadsheets: the purchase register, and GSTR-2B when it comes as Excel
# ---------------------------------------------------------------------------

#: Header words each field goes by in the exports firms actually have (Tally,
#: the portal's Excel, hand-made sheets). Compared after lower-casing and
#: dropping everything that is not a letter or digit.
_HEADERS = {
    "gstin": ["gstin", "gstinuin", "gstinofsupplier", "supplierGSTIN", "gstinuin", "ctin", "gstno", "partygstin"],
    "invoice_no": ["supplierinvoiceno", "suppliersinvoiceno", "partyinvoiceno", "invoiceno", "invoicenumber", "invno", "billno", "billnumber", "docno", "invoice", "vchno", "voucherno"],
    "invoice_date": ["supplierinvoicedate", "partyinvoicedate", "invoicedate", "date", "billdate", "invdate", "docdate", "vchdate"],
    "taxable": ["taxablevalue", "value", "grossvalue", "taxable", "taxableamount", "assessablevalue", "basicamount"],
    "igst": ["igst", "integratedtax", "igstamount", "integratedtaxamount"],
    "cgst": ["cgst", "centraltax", "cgstamount", "centraltaxamount"],
    "sgst": ["sgst", "sgstutgst", "statetax", "sgstamount", "stateuttax", "statetaxamount"],
    "cess": ["cess", "cessamount"],
    "supplier": ["tradename", "suppliername", "partyname", "party", "vendor", "vendorname", "name", "tradelegalname", "particulars"],
    "hsn": ["hsn", "hsncode", "sac", "hsnsac", "hsnsaccode"],
    "category": ["category", "ledger", "ledgername", "nature", "natureofexpense", "expensehead"],
    "rcm": ["rcm", "reversecharge", "reversechargeapplicable"],
    "type": ["type", "documenttype", "doctype", "invoicetype"],
}
_HEADERS = {k: [re.sub(r"[^a-z0-9]", "", w.lower()) for w in v] for k, v in _HEADERS.items()}

REQUIRED = ("gstin", "invoice_no")


def _key(text) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def _read_table(data: bytes, filename: str) -> list[list]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook  # local: only spreadsheet uploads need it

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises a zoo of types on a bad file
            raise GstParseError(f"Cannot open this Excel file: {exc}") from exc
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith((".csv", ".txt")):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        return [list(r) for r in csv.reader(io.StringIO(text))]
    raise GstParseError("Upload an .xlsx or .csv file.")


#: A tax head that has no column of its own is summed from every column whose
#: header names it -- Tally writes one per ledger and rate. Rate columns and
#: totals are not amounts and are left out.
_TAX_TOKENS = {"igst": "igst", "cgst": "cgst", "sgst": "sgst", "cess": "cess"}
_NOT_AMOUNT = ("rate", "total", "tax%", "percent")


def _tax_columns(keys: list[str], head: str) -> list[int]:
    token = _TAX_TOKENS[head]
    out = []
    for i, k in enumerate(keys):
        if token not in k or any(w in k for w in _NOT_AMOUNT):
            continue
        # "sgst" also appears inside "sgstutgst"; "igst" is not inside "cgst".
        out.append(i)
    return out


def detect_columns(header: list, mapping: dict[str, str] | None = None) -> dict[str, list[int]]:
    """Field -> the column(s) it is read from. ``mapping`` (field -> header) overrides."""
    keys = [_key(h) for h in header]
    found: dict[str, list[int]] = {}
    for field_, words in _HEADERS.items():
        override = (mapping or {}).get(field_)
        if override is not None:
            if _key(override) in keys:
                found[field_] = [keys.index(_key(override))]
            continue
        for w in words:
            if w in keys:
                found[field_] = [keys.index(w)]
                break
        else:
            if field_ in _TAX_TOKENS:
                cols = _tax_columns(keys, field_)
                if cols:
                    found[field_] = cols
    return found


def _cell_reader(row: list, columns: dict[str, list[int]]):
    def cell(field_):
        idx = columns.get(field_)
        if not idx or idx[0] >= len(row):
            return None
        return row[idx[0]]

    return cell


def _amount(row: list, columns: dict[str, list[int]], field_: str, where: str) -> int:
    """A money field, summing across columns when a head is split by rate."""
    total = 0
    for i in columns.get(field_, []):
        if i < len(row):
            total += _paise(row[i], f"{where} {field_}")
    return total


def parse_register(
    data: bytes, filename: str, *, mapping: dict[str, str] | None = None, portal: bool = False
) -> list[Invoice]:
    """A purchase register (or a portal Excel) as invoices.

    Finds the header row as the first row that names an invoice-number column,
    since exports put titles and company names above it.
    """
    if mapping is not None and not (
        isinstance(mapping, dict) and all(isinstance(k, str) and isinstance(v, str) for k, v in mapping.items())
    ):
        raise GstParseError("The column mapping should be an object of field names to header names.")
    rows = _read_table(data, filename)
    header_at = None
    for i, row in enumerate(rows[:30]):
        cols = detect_columns(row, mapping)
        if "invoice_no" in cols:
            header_at, columns = i, cols
            break
    if header_at is None:
        raise GstParseError(
            "Could not find a header row with an invoice number column. "
            "Expected something like 'Invoice No', 'GSTIN', 'Taxable Value', 'CGST'."
        )
    missing = [f for f in REQUIRED if f not in columns]
    if missing:
        raise GstParseError(f"No column found for: {', '.join(missing)}.")
    if not any(f in columns for f in ("igst", "cgst", "sgst")):
        raise GstParseError("No tax columns found (IGST / CGST / SGST).")

    out: list[Invoice] = []
    for offset, row in enumerate(rows[header_at + 1 :], start=header_at + 2):
        if not any(c not in (None, "") for c in row):
            continue

        cell = _cell_reader(row, columns)

        where = f"row {offset}"
        number = str(cell("invoice_no") or "").strip()
        if not number:
            continue  # totals and blank lines carry no invoice number
        doc_type = str(cell("type") or "").lower()
        out.append(
            Invoice(
                ref=f"{'2b' if portal else 'reg'}:{offset}",
                gstin=normalise_gstin(str(cell("gstin") or "")),
                invoice_no=number,
                invoice_date=_date(cell("invoice_date"), where),
                taxable_paise=_amount(row, columns, "taxable", where),
                igst_paise=_amount(row, columns, "igst", where),
                cgst_paise=_amount(row, columns, "cgst", where),
                sgst_paise=_amount(row, columns, "sgst", where),
                cess_paise=_amount(row, columns, "cess", where),
                supplier_name=str(cell("supplier") or "").strip(),
                hsn=str(cell("hsn") or "").strip(),
                category=str(cell("category") or "").strip(),
                section=Section.CDN if ("credit" in doc_type or "debit" in doc_type) else Section.B2B,
                rcm=_truthy(cell("rcm")),
            )
        )
    if not out:
        raise GstParseError("The file has a header but no invoice rows.")
    return out
