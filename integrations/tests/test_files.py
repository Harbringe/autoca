"""Any file a client sends becomes the document shape the readers know: what it is comes from its bytes, never its name."""

from __future__ import annotations

import datetime
import io
import zipfile

import pytest

from banking.parsers import parse_statement
from integrations import files
from integrations.pdf.base import PdfExtractionError


def workbook(rows, sheets=1) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    for index in range(sheets):
        sheet = book.active if index == 0 else book.create_sheet()
        for row in rows:
            sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def docx(paragraphs, table=None) -> bytes:
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    body = "".join(f"<w:p><w:r><w:t>{p}</w:t></w:r></w:p>" for p in paragraphs)
    if table:
        rows = "".join(
            "<w:tr>" + "".join(f"<w:tc><w:p><w:r><w:t>{c}</w:t></w:r></w:p></w:tc>" for c in row) + "</w:tr>" for row in table
        )
        body += f"<w:tbl>{rows}</w:tbl>"
    xml = f'<?xml version="1.0"?><w:document xmlns:w="{w}"><w:body>{body}</w:body></w:document>'
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("word/document.xml", xml)
    return out.getvalue()


def png(width=40, height=30) -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(out, format="PNG")
    return out.getvalue()


STATEMENT = [
    ["HDFC BANK LTD"],
    ["Statement of account"],
    ["Account No :", "50100123456789"],
    ["IFSC :", "HDFC0000123"],
    ["From :", "01/04/2025", "To :", "30/04/2025"],
    [],
    ["Date", "Narration", "Chq/Ref No", "Withdrawal Amt.", "Deposit Amt.", "Closing Balance"],
    [datetime.datetime(2025, 4, 1), "OPENING BALANCE", None, None, None, 25000.0],
    [datetime.datetime(2025, 4, 3), "UPI-SWIGGY-9876543210", "000000123456", 450.0, None, 24550.0],
    [datetime.datetime(2025, 4, 7), "NEFT CR-SBIN0001234-ACME PVT LTD", "N123456789", None, 120000.0, 144550.0],
    [datetime.datetime(2025, 4, 12), "ATW-4321-CASH WDL", None, 10000.0, None, 134550.0],
    [datetime.datetime(2025, 4, 30), "CLOSING BALANCE", None, None, None, 134550.0],
]


def test_what_a_file_is_comes_from_its_bytes_not_its_name():
    assert files.sniff(b"%PDF-1.7 ...", "x.xlsx") == files.PDF
    assert files.sniff(workbook([["a"]]), "statement.pdf") == files.XLSX
    assert files.sniff(docx(["hello"]), "statement.pdf") == files.DOCX
    assert files.sniff(png(), "statement.pdf") == files.IMAGE
    assert files.sniff(b"Date,Amount\n1,2\n", "s.csv") == files.CSV
    assert files.sniff(b"MZ\x90\x00 an executable", "statement.pdf") is None
    assert files.sniff(b"Date,Amount\n1,2\n", "s.exe") is None


def test_an_old_office_file_or_a_heic_photo_says_what_to_do_instead():
    assert "save it as .xlsx" in files.describe_refusal(b"\xd0\xcf\x11\xe0....", "old.xls")
    assert "JPG or PNG" in files.describe_refusal(b"....", "photo.heic")


def test_an_excel_statement_is_read_and_proved_like_a_pdf_one():
    loaded = files.load(workbook(STATEMENT), "statement.xlsx")

    statement = parse_statement(loaded.document)

    assert loaded.kind == files.XLSX and loaded.images is None
    assert len(statement) == 3 and statement.account_number.endswith("56789")
    assert statement.opening_balance_paise == 25_000_00 and statement.closing_balance_paise == 134_550_00
    assert statement.transactions[0].date == datetime.date(2025, 4, 3) and statement.transactions[0].debit_paise == 450_00


def test_a_wrong_figure_in_the_sheet_still_breaks_the_balance_proof():
    bad = [list(r) for r in STATEMENT]
    bad[9][5] = 144000.0  # the running balance no longer follows

    with pytest.raises(Exception, match="(?i)balance|add up|tie"):
        parse_statement(files.load(workbook(bad), "s.xlsx").document)


def test_a_csv_statement_is_read_the_same_way():
    lines = [",".join(f'"{c}"' if c is not None else "" for c in ["" if x is None else (x.strftime("%d-%m-%Y") if isinstance(x, datetime.datetime) else x) for x in row]) for row in STATEMENT]
    text = "\n".join(lines).encode()

    statement = parse_statement(files.load(text, "statement.csv").document)

    assert len(statement) == 3 and statement.closing_balance_paise == 134_550_00


def test_a_word_file_gives_its_text_and_its_tables():
    loaded = files.load(docx(["Tax Invoice", "Invoice No: RT/42"], table=[["Item", "Amount"], ["Goods", "1,000.00"]]), "inv.docx")

    page = loaded.document.pages[0]
    assert "Invoice No: RT/42" in page.text and "Goods  1,000.00" in page.text
    assert page.tables[0][1] == ("Goods", "1,000.00")


def test_a_photo_has_no_text_and_comes_with_its_page_images():
    loaded = files.load(png(), "scan.png")

    assert loaded.is_image and not loaded.document.has_text_layer
    assert loaded.document.page_count == 1 and len(loaded.images) == 1 and loaded.images[0].startswith(b"\x89PNG")


def test_a_big_photo_is_scaled_down_before_it_is_sent_anywhere():
    from PIL import Image

    image = Image.open(io.BytesIO(files.load(png(5000, 3000), "big.png").images[0]))

    assert max(image.size) == files.MAX_IMAGE_SIDE


def test_a_zip_that_claims_to_be_enormous_is_refused_before_it_is_opened(monkeypatch):
    monkeypatch.setattr(files, "MAX_UNCOMPRESSED_BYTES", 10)

    with pytest.raises(PdfExtractionError, match="larger than it looks"):
        files.load(workbook([["a", "b"]]), "s.xlsx")


def test_an_empty_workbook_is_one_empty_page_not_a_crash():
    loaded = files.load(workbook([[]]), "s.xlsx")

    assert loaded.document.page_count == 1 and not loaded.document.has_text_layer


def test_something_that_is_none_of_these_is_refused_in_words():
    with pytest.raises(files.UnsupportedFileError, match="cannot be read"):
        files.load(b"MZ\x90\x00 an executable", "statement.pdf")
