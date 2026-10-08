"""The financial statements as an Excel workbook, laid out the way the ICAI format sets them out.

Sheets: Balance Sheet, Statement of P&L, Notes 1 to 3 (about the entity, policies, the partner-wise capital table) and
Notes 4 to 25 with each note's sub-heads. Figures are in the unit the client chose (rupees, lakhs, ...), already rounded by
``ledger.nce``; every text that came from a person or a ledger name goes through ``_safe_cell`` so a name beginning with
``=`` cannot become a formula.
"""

from __future__ import annotations

import io
from decimal import Decimal

from ledger.nce import Statements

_HEAD_FILL = "D9E2F3"


def _figure(paise, unit: int):
    if paise is None:
        return None
    return float(Decimal(paise) / Decimal(unit))


def workbook(client, financial_year: int, s: Statements) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    from gst.report import (
        _safe_cell,  # a function-level import: gst.services reaches back into ledger
    )

    unit = s.unit_paise
    number_format = '#,##0.00;(#,##0.00);"-"'
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    head = PatternFill("solid", fgColor=_HEAD_FILL)
    line = Border(bottom=Side(style="thin"))
    end, prev_end = f"31 March {financial_year + 1}", f"31 March {financial_year}"

    wb = Workbook()

    def sheet(title: str, subtitle: str, widths: list[int], first: bool = False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append([_safe_cell(client.name)])
        ws.append([subtitle])
        ws.append([f"(Amount in {s.unit_label})"])
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"].font = bold
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        return ws

    def statement(title: str, subtitle: str, rows, first: bool = False):
        ws = sheet(title, subtitle, [64, 8, 20, 20], first)
        ws.append([])
        ws.append(["Particulars", "Note", end, prev_end if s.has_previous else ""])
        for c in ws[ws.max_row]:
            c.font = bold
            c.fill = head
        for r in rows:
            heading = r.kind == "heading"
            ws.append(
                [
                    ("    " * r.level) + _safe_cell(r.label),
                    r.note,
                    None if heading else _figure(r.current_paise, unit),
                    None if heading or not s.has_previous else _figure(r.previous_paise, unit),
                ]
            )
            row = ws[ws.max_row]
            row[0].alignment = wrap
            row[2].number_format = row[3].number_format = number_format
            if heading or r.kind == "total":
                for c in row:
                    c.font = bold
            if r.kind == "total":
                for c in row:
                    c.border = line
        return ws

    statement("Balance Sheet", f"Balance Sheet as at {end}", s.balance_sheet, first=True)
    statement("Statement of P&L", f"Statement of Profit and Loss for the year ended {end}", s.profit_and_loss)

    ws = sheet("Notes 1 to 3", f"Notes forming part of the Financial Statements for the year ended {end}", [8, 34, 14, 18, 18, 18, 18, 18, 18, 18])
    ws.append([])
    ws.append(["Note 1", "Brief about the entity"])
    ws["A5"].font = ws["B5"].font = bold
    ws.append(["", _safe_cell(s.about) or "Not written."])
    ws.append(["Note 2", "Significant Accounting Policies"])
    ws[f"A{ws.max_row}"].font = ws[f"B{ws.max_row}"].font = bold
    ws.append(["", ((_safe_cell(s.policies) + "\n\n") if s.policies else "") + s.size_statement])
    for r in (6, 8):
        ws[f"B{r}"].alignment = wrap
        ws.merge_cells(f"B{r}:J{r}")
        ws.row_dimensions[r].height = 90
    ws.append([])
    ws.append(["Note 3", s.capital_title])
    ws[f"A{ws.max_row}"].font = ws[f"B{ws.max_row}"].font = bold
    ws.append(
        ["Sr.", "Name of Partner/ Proprietor/ Owner", "Share of profit/ (loss) (%)", "Opening balance", "Capital introduced", "Remuneration", "Interest", "Withdrawals", "Share of profit/ (loss)", "Closing balance"]
    )
    for c in ws[ws.max_row]:
        c.font = bold
        c.fill = head
        c.alignment = wrap
    if s.capital and s.capital.rows:
        for i, p in enumerate(s.capital.rows, 1):
            ws.append(
                [i, _safe_cell(p.name), p.share_bp / 100]
                + [_figure(v, unit) for v in (p.opening_paise, p.introduced_paise, p.remuneration_paise, p.interest_paise, p.withdrawals_paise, p.profit_share_paise, p.closing_paise)]
            )
        ws.append(
            ["", "Total", sum(p.share_bp for p in s.capital.rows) / 100]
            + [_figure(s.capital.total(f), unit) for f in ("opening_paise", "introduced_paise", "remuneration_paise", "interest_paise", "withdrawals_paise", "profit_share_paise", "closing_paise")]
        )
        for c in ws[ws.max_row]:
            c.font = bold
        if s.capital.previous:
            ws.append(
                ["", "Previous year"]
                + [""]
                + [_figure(s.capital.total(f, True), unit) for f in ("opening_paise", "introduced_paise", "remuneration_paise", "interest_paise", "withdrawals_paise", "profit_share_paise", "closing_paise")]
            )
    else:
        ws.append(["", "Partners have not been entered yet."])
    for row in ws.iter_rows(min_row=1, min_col=4, max_col=10):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = number_format

    ws = sheet("Notes 4 to 25", f"Notes forming part of the Financial Statements for the year ended {end}", [8, 70, 20, 20])
    for note in s.notes:
        ws.append([])
        ws.append([note.number, note.title, end, prev_end if s.has_previous else ""])
        for c in ws[ws.max_row]:
            c.font = bold
            c.fill = head
        section = None
        for r in note.rows:
            if r.section and r.section != section:
                section = r.section
                ws.append(["", _safe_cell(section)])
                ws[f"B{ws.max_row}"].font = Font(bold=True, italic=True)
            ws.append(["", ("    " if r.section else "") + _safe_cell(r.label), _figure(r.current_paise, unit), _figure(r.previous_paise, unit) if s.has_previous else None])
            ws[f"C{ws.max_row}"].number_format = ws[f"D{ws.max_row}"].number_format = number_format
        ws.append(["", "Total", _figure(note.total_current_paise, unit), _figure(note.total_previous_paise, unit) if s.has_previous else None])
        for c in ws[ws.max_row]:
            c.font = bold
            c.border = line
        ws[f"C{ws.max_row}"].number_format = ws[f"D{ws.max_row}"].number_format = number_format
        for sched in (x for x in s.schedules if x.note == note.number):
            ws.append([])
            ws.append(["", _safe_cell(sched.title), *sched.columns])
            for c in ws[ws.max_row]:
                c.font = bold
            for r in sched.rows:
                ws.append(["", _safe_cell(r.label), *[None if v is None else _figure(v, unit) for v in r.values]])
                for c in ws[ws.max_row]:
                    c.font = Font(bold=r.kind != "line")
                    c.number_format = number_format

    if s.regroupings:
        ws.append([])
        ws.append(["", "Regrouping of previous year figures"])
        ws[f"B{ws.max_row}"].font = bold
        for g in s.regroupings:
            ws.append(["", _safe_cell(g.text)])
            ws[f"B{ws.max_row}"].alignment = wrap

    if s.warnings:
        ws.append([])
        ws.append(["", "To settle before issuing"])
        ws[f"B{ws.max_row}"].font = bold
        for w in s.warnings:
            ws.append(["", w])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
