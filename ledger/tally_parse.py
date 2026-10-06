"""Reading a Tally masters export: ledgers, their groups, and their opening balances.

Nothing here touches the database. It turns the bytes a person uploaded into
:class:`ParsedMasters`, and the rest of the import (``ledger.tally_import``)
decides what to do with them. Keeping the two apart means the part that faces
an untrusted file can be tested, and attacked, without a tenant in sight.

**The file is hostile until proved otherwise.** It is XML from a program the
firm does not control, or a spreadsheet somebody edited by hand.

* XML goes through ``defusedxml`` with DTDs, entities and external references
  refused, so an entity bomb or an external fetch is a refusal, not an event.
* The type is decided from the bytes, not the extension, and the size, the
  number of ledgers and the length of every name are capped.
* Every string is a name typed by a stranger: whitespace is collapsed, control
  characters are refused, and a name a spreadsheet would run as a formula
  (``= + - @``) is skipped with a reason rather than carried into a later export.
* Any problem with one row skips that row, with the reason, and the file goes on.
  Any problem with the file as a whole is a :class:`TallyParseError` whose message
  is written for the person who uploaded it. Nothing else escapes.

**Sign convention.** Tally writes a debit as a negative amount and a credit as
a positive one (and ``Dr``/``Cr`` suffixes in spreadsheets). ``core.money.to_paise``
reads ``Dr`` as negative to match, so what it returns is Tally's sign. The rest
of this system is debit-positive, so every amount is negated once, here, and
leaves this module in the books' own convention.

Element names, group names and header names below that come from memory of
Tally rather than from a real file are marked ``UNVERIFIED``.
"""

from __future__ import annotations

import csv
import datetime
import io
import logging
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field

from classify.models import LedgerGroup
from core.money import MAX_PAISE, to_paise

logger = logging.getLogger("autoca.tally")

MAX_LEDGERS = 20_000
MAX_GROUPS = 20_000
MAX_NAME_LENGTH = 255
MAX_PATH_LENGTH = 512
MAX_GROUP_DEPTH = 24
MAX_SHEET_ROWS = 100_000
#: What a workbook may expand to. A spreadsheet is a zip, and a zip can be a bomb.
MAX_XLSX_UNPACKED_BYTES = 100 * 1024 * 1024
MAX_XLSX_MEMBERS = 2_000

ALLOWED_EXTENSIONS = (".xml", ".xlsx", ".csv")


class TallyParseError(ValueError):
    """The file as a whole cannot be read. The message is for the person who uploaded it."""


class TallyTooLargeError(ValueError):
    """The file is bigger than an import will take."""


@dataclass(frozen=True)
class MasterGroup:
    name: str
    parent: str


@dataclass(frozen=True)
class MasterLedger:
    #: Position in the file (an element count in XML, the sheet row in a spreadsheet).
    ordinal: int
    name: str
    #: The group the file puts it under, as written. May be empty.
    parent: str
    #: Debits positive, like the books. Zero when the file gives none.
    opening_paise: int
    alias: str | None = None


@dataclass(frozen=True)
class SkippedLine:
    ordinal: int
    #: The name when it could be read, so the person can find the line. Otherwise empty.
    name: str
    reason: str


@dataclass
class ParsedMasters:
    source_format: str
    ledgers: list[MasterLedger] = field(default_factory=list)
    groups: dict[str, MasterGroup] = field(default_factory=dict)
    skipped: list[SkippedLine] = field(default_factory=list)
    #: The date the company's books begin from, when the file says (XML only).
    books_from: datetime.date | None = None


# ---------------------------------------------------------------------------
# Names
# ---------------------------------------------------------------------------

_FORMULA_START = ("=", "+", "-", "@")


class _BadText(ValueError):
    """A string that cannot be used. The message is the reason shown against the line."""


def clean_text(raw, *, limit: int = MAX_NAME_LENGTH, what: str = "name") -> str:
    """One name, with whitespace collapsed. Raises :class:`_BadText` with a reason."""
    text = " ".join(str(raw if raw is not None else "").split())
    text = unicodedata.normalize("NFC", text)
    if not text:
        raise _BadText(f"The {what} is empty.")
    if len(text) > limit:
        raise _BadText(f"The {what} is longer than {limit} characters.")
    if any(unicodedata.category(c) in ("Cc", "Cf", "Cs") for c in text):
        raise _BadText(f"The {what} contains control characters.")
    if text.startswith(_FORMULA_START):
        raise _BadText(
            f"The {what} starts with {text[0]!r}, which a spreadsheet would run as a formula. "
            f"Rename it in Tally and export again."
        )
    return text


def normal_name(name: str) -> str:
    """What two spellings of one ledger have in common: NFKC, case-folded, single-spaced."""
    return " ".join(unicodedata.normalize("NFKC", name).casefold().split())


# ---------------------------------------------------------------------------
# Groups
# ---------------------------------------------------------------------------

G = LedgerGroup

#: Tally's predefined groups -> ours. Keyed by ``normal_name``. Predefined groups are
#: not in an export, so a ledger's chain of parents ends at one of these names.
#: UNVERIFIED: needs a real sample (exact spellings, e.g. "Bank OCC A/c", "Branch / Divisions").
TALLY_GROUPS: dict[str, str] = {
    "sundry debtors": G.DEBTOR,
    "sundry creditors": G.CREDITOR,
    "bank accounts": G.BANK,
    "bank od a/c": G.BANK_OD,
    "bank occ a/c": G.BANK_OD,
    "cash-in-hand": G.CASH,
    "duties & taxes": G.DUTIES_AND_TAXES,
    "direct expenses": G.DIRECT_EXPENSE,
    "expenses (direct)": G.DIRECT_EXPENSE,
    "indirect expenses": G.INDIRECT_EXPENSE,
    "expenses (indirect)": G.INDIRECT_EXPENSE,
    "direct incomes": G.DIRECT_INCOME,
    "income (direct)": G.DIRECT_INCOME,
    "indirect incomes": G.INDIRECT_INCOME,
    "income (indirect)": G.INDIRECT_INCOME,
    "capital account": G.CAPITAL,
    "loans (liability)": G.LOAN,
    "secured loans": G.LOAN,
    "unsecured loans": G.LOAN,
    "investments": G.INVESTMENT,
    "suspense a/c": G.SUSPENSE,
    "fixed assets": G.FIXED_ASSET,
    "stock-in-hand": G.STOCK,
    "current assets": G.CURRENT_ASSET,
    "loans & advances (asset)": G.LOAN_ADVANCE,
    "deposits (asset)": G.DEPOSIT,
    "misc. expenses (asset)": G.MISC_EXPENDITURE,
    "current liabilities": G.CURRENT_LIABILITY,
    "provisions": G.PROVISION,
    "reserves & surplus": G.RESERVES,
    "retained earnings": G.RESERVES,
    "sales accounts": G.SALES,
    "purchase accounts": G.PURCHASE,
}


@dataclass(frozen=True)
class GroupResolution:
    #: One of ``LedgerGroup`` values, or None when the chain ends at a group of the
    #: firm's own that we cannot place.
    group: str | None
    #: The ledger's group, then each parent up to the one that decided it.
    path: tuple[str, ...]


def resolve_group(parent: str, groups: dict[str, MasterGroup]) -> GroupResolution:
    """Walk up from ``parent`` to the nearest group we know.

    The *nearest* known ancestor, not the primary one: a ledger under Sundry
    Debtors is a debtor, though Sundry Debtors sits under Current Assets. A cycle
    or an excessive depth ends the walk unplaced rather than looping.
    """
    path: list[str] = []
    seen: set[str] = set()
    name = parent
    while name and len(path) < MAX_GROUP_DEPTH:
        key = normal_name(name)
        if key in seen:
            break
        seen.add(key)
        path.append(name)
        if key in TALLY_GROUPS:
            return GroupResolution(TALLY_GROUPS[key], tuple(path))
        node = groups.get(key)
        name = node.parent if node else ""
    return GroupResolution(None, tuple(path))


def path_text(path: tuple[str, ...]) -> str | None:
    text = " > ".join(path)
    return text[:MAX_PATH_LENGTH] if text else None


# ---------------------------------------------------------------------------
# Amounts
# ---------------------------------------------------------------------------

_AMOUNT = re.compile(r"^\(?[-+]?\d[\d,]*(\.\d+)?\)?\s*(dr|cr)?$", re.IGNORECASE)


class _BadAmount(ValueError):
    pass


def parse_opening(text) -> int:
    """An opening balance as debit-positive paise. ``''`` is zero.

    Accepts Tally's signed XML amounts (``-1500.00`` is a debit) and the ``Dr`` /
    ``Cr`` form of its spreadsheets. Refuses more than two decimals, exponents,
    currency words and anything else it cannot be sure of.
    """
    cleaned = " ".join(str(text if text is not None else "").replace("−", "-").split())
    if not cleaned:
        return 0
    if not _AMOUNT.match(cleaned):
        raise _BadAmount(f"The opening balance {cleaned[:40]!r} is not an amount.")
    try:
        tally_signed = to_paise(cleaned)
    except (ValueError, ArithmeticError) as exc:
        raise _BadAmount(f"The opening balance {cleaned[:40]!r} is not an exact amount in rupees and paise.") from exc
    if abs(tally_signed) >= MAX_PAISE:
        raise _BadAmount("The opening balance is too large to be real.")
    return -tally_signed


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def parse_masters(data: bytes, filename: str, *, max_bytes: int) -> ParsedMasters:
    """Read a masters export. ``max_bytes`` is the caller's size limit."""
    if len(data) > max_bytes:
        raise TallyTooLargeError(
            f"This file is {len(data) // (1024 * 1024)} MB; an import takes up to "
            f"{max_bytes // (1024 * 1024)} MB. Export only the ledger masters and try again."
        )
    if not data.strip(b"\x00 \r\n\t"):
        raise TallyParseError("The file is empty.")

    extension = _extension(filename)
    if extension not in ALLOWED_EXTENSIONS:
        raise TallyParseError("Upload the masters as an .xml, .xlsx or .csv file exported from Tally.")

    kind = _sniff(data)
    if kind != extension[1:]:
        raise TallyParseError(
            f"This file's contents are not {extension[1:].upper()}, whatever its name says. "
            f"Export the masters from Tally again and upload that file."
        )

    try:
        if kind == "xml":
            parsed = _parse_xml(data)
        elif kind == "xlsx":
            parsed = _parse_xlsx(data)
        else:
            parsed = _parse_csv(data)
    except (TallyParseError, TallyTooLargeError):
        raise
    except Exception as exc:  # noqa: BLE001 -- a hostile file must end as a refusal, never a 500
        logger.warning("tally file refused: %s", type(exc).__name__)
        raise TallyParseError(
            "This file could not be read as a Tally masters export. Export the ledger masters "
            "from Tally again and upload that file."
        ) from exc

    if not parsed.ledgers and not parsed.skipped:
        raise TallyParseError(
            "No ledgers were found in this file. In Tally, export Masters (ledgers) as XML, "
            "or the List of Accounts as Excel."
        )
    return parsed


def _extension(filename: str) -> str:
    name = (filename or "").lower()
    dot = name.rfind(".")
    return name[dot:] if dot >= 0 else ""


def _sniff(data: bytes) -> str:
    if data.startswith(b"PK\x03\x04"):
        return "xlsx"
    text = _decode(data)
    return "xml" if text.lstrip().startswith("<") else "csv"


# ---------------------------------------------------------------------------
# Text decoding
# ---------------------------------------------------------------------------


def _decode(data: bytes) -> str:
    """UTF-8 or UTF-16 (Tally's usual), with or without a byte-order mark."""
    if data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff"):
        codecs = ("utf-16",)
    elif data.startswith(b"\xef\xbb\xbf"):
        codecs = ("utf-8-sig",)
    elif data[:2] == b"<\x00":
        codecs = ("utf-16-le",)
    elif data[:2] == b"\x00<":
        codecs = ("utf-16-be",)
    else:
        codecs = ("utf-8", "cp1252")
    for codec in codecs:
        try:
            text = data.decode(codec)
        except UnicodeDecodeError:
            continue
        if "\x00" in text:
            break
        return text
    raise TallyParseError(
        "This file's text encoding could not be read. Export it from Tally again (UTF-8 or UTF-16)."
    )


# ---------------------------------------------------------------------------
# XML
# ---------------------------------------------------------------------------

_XML_DECLARATION = re.compile(r"^\s*<\?xml[^>]*\?>")
_CHAR_REF = re.compile(r"&#(?:[xX]([0-9a-fA-F]{1,8})|([0-9]{1,9}));")
_ILLEGAL_RAW = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")


def _illegal_code(code: int) -> bool:
    return (
        code == 0
        or (code < 0x20 and code not in (0x09, 0x0A, 0x0D))
        or 0xD800 <= code <= 0xDFFF
        or code in (0xFFFE, 0xFFFF)
        or code > 0x10FFFF
    )


def _strip_illegal_xml(text: str) -> str:
    """Remove what XML 1.0 forbids, which Tally writes anyway -- raw, or as ``&#4;``.

    UNVERIFIED: needs a real sample. Tally is remembered to put ``&#4;`` in front of its
    reserved names such as Primary; a strict parser rejects it outright.
    """

    def drop(match: re.Match) -> str:
        code = int(match.group(1), 16) if match.group(1) else int(match.group(2))
        return "" if _illegal_code(code) else match.group(0)

    return _ILLEGAL_RAW.sub("", _CHAR_REF.sub(drop, text))


def _local(tag) -> str:
    return tag.rsplit("}", 1)[-1].upper() if isinstance(tag, str) else ""


def _parse_xml(data: bytes) -> ParsedMasters:
    from defusedxml import DefusedXmlException
    from defusedxml.ElementTree import ParseError, iterparse

    text = _strip_illegal_xml(_decode(data))
    # Decoded already, so a declaration naming UTF-16 would now be a lie.
    text = _XML_DECLARATION.sub("", text, count=1)
    source = io.BytesIO(text.encode("utf-8"))

    parsed = ParsedMasters(source_format="xml")
    ordinal = 0
    try:
        for _event, element in iterparse(
            source, events=("end",), forbid_dtd=True, forbid_entities=True, forbid_external=True
        ):
            tag = _local(element.tag)
            # UNVERIFIED: needs a real sample. LEDGER / GROUP / BOOKSFROM element names.
            if tag == "LEDGER":
                ordinal += 1
                if ordinal > MAX_LEDGERS:
                    raise TallyParseError(
                        f"This file has more than {MAX_LEDGERS:,} ledgers, which is more than an "
                        f"import takes. Export the masters in parts."
                    )
                _read_ledger(element, ordinal, parsed)
                element.clear()
            elif tag == "GROUP":
                if len(parsed.groups) >= MAX_GROUPS:
                    raise TallyParseError(f"This file has more than {MAX_GROUPS:,} groups.")
                _read_group(element, parsed)
                element.clear()
            elif tag == "BOOKSFROM" and parsed.books_from is None:
                parsed.books_from = _tally_date(element.text)
    except DefusedXmlException as exc:
        raise TallyParseError(
            "This file contains definitions (a DTD or entities) that an export from Tally never "
            "has, so it was not read."
        ) from exc
    except ParseError as exc:
        raise TallyParseError(
            "This file is not well-formed XML, so it cannot be read. Export the masters from Tally again."
        ) from exc
    return parsed


def _tally_date(raw) -> datetime.date | None:
    text = (raw or "").strip()
    try:
        return datetime.datetime.strptime(text, "%Y%m%d").date()
    except ValueError:
        return None


def _element_name(element) -> tuple[str, list[str]]:
    """The name of a LEDGER or GROUP, and the other names it answers to.

    UNVERIFIED: needs a real sample. NAME attribute, ``LANGUAGENAME.LIST/NAME.LIST/NAME``
    (the first is the name, the rest aliases), then a bare NAME child.
    """
    names = [
        (n.text or "").strip()
        for n in element.findall("LANGUAGENAME.LIST/NAME.LIST/NAME")
        if (n.text or "").strip()
    ]
    primary = (element.get("NAME") or "").strip() or (names[0] if names else "") or (
        (element.findtext("NAME") or "").strip()
    )
    return primary, [n for n in names if normal_name(n) != normal_name(primary)]


def _is_deleted(element) -> bool:
    # UNVERIFIED: needs a real sample. ACTION="Delete" on a master.
    return (element.get("ACTION") or "").strip().lower() == "delete"


def _read_group(element, parsed: ParsedMasters) -> None:
    if _is_deleted(element):
        return
    raw_name, _ = _element_name(element)
    raw_parent = (element.findtext("PARENT") or "").strip()
    try:
        name = clean_text(raw_name, what="group name")
        parent = clean_text(raw_parent, what="group name") if raw_parent else ""
    except _BadText:
        return
    parsed.groups[normal_name(name)] = MasterGroup(name=name, parent=parent)


def _read_ledger(element, ordinal: int, parsed: ParsedMasters) -> None:
    raw_name, aliases = _element_name(element)
    if _is_deleted(element):
        parsed.skipped.append(SkippedLine(ordinal, _shown(raw_name), "Marked as deleted in Tally."))
        return
    try:
        name = clean_text(raw_name)
    except _BadText as exc:
        parsed.skipped.append(SkippedLine(ordinal, _shown(raw_name), str(exc)))
        return

    parent = ""
    raw_parent = (element.findtext("PARENT") or "").strip()
    if raw_parent:
        try:
            parent = clean_text(raw_parent, what="group name")
        except _BadText as exc:
            parsed.skipped.append(SkippedLine(ordinal, name, str(exc)))
            return

    try:
        opening = parse_opening(element.findtext("OPENINGBALANCE"))  # UNVERIFIED: needs a real sample
    except _BadAmount as exc:
        parsed.skipped.append(SkippedLine(ordinal, name, str(exc)))
        return

    alias = None
    for candidate in aliases:
        try:
            alias = clean_text(candidate, what="alias")
            break
        except _BadText:
            continue
    parsed.ledgers.append(MasterLedger(ordinal, name, parent, opening, alias))


def _shown(raw) -> str:
    """A name safe to put in a message even when it failed the checks."""
    text = " ".join(str(raw or "").split())
    return "".join(c for c in text if unicodedata.category(c) not in ("Cc", "Cf", "Cs"))[:80]


# ---------------------------------------------------------------------------
# Spreadsheets
# ---------------------------------------------------------------------------

#: UNVERIFIED: needs a real sample. Column headings of Tally's List of Accounts export.
_HEADERS = {
    "name": {"name", "ledger", "ledger name", "particulars"},
    "group": {"group", "group name", "under", "parent"},
    "opening": {"opening", "opening balance", "opening bal", "opening bal.", "opening balance (dr/cr)"},
    "side": {"dr/cr", "drcr", "dr cr", "opening dr/cr", "opening drcr", "type"},
    "debit": {"debit", "opening debit", "dr"},
    "credit": {"credit", "opening credit", "cr"},
}


def _parse_csv(data: bytes) -> ParsedMasters:
    text = _decode(data)
    sample = text[:4096]
    delimiter = max(",;\t|", key=sample.count)
    try:
        rows = [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    except csv.Error as exc:
        raise TallyParseError("This CSV file could not be read. Export the masters from Tally again.") from exc
    if len(rows) > MAX_SHEET_ROWS:
        raise TallyParseError(f"This file has more than {MAX_SHEET_ROWS:,} rows.")
    return _read_table(rows, "csv")


def _parse_xlsx(data: bytes) -> ParsedMasters:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise TallyParseError("This is not a valid Excel workbook.") from exc
    infos = archive.infolist()
    if len(infos) > MAX_XLSX_MEMBERS or sum(i.file_size for i in infos) > MAX_XLSX_UNPACKED_BYTES:
        raise TallyParseError("This workbook is too large once unpacked to be a list of ledgers.")

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True, max_col=200):
                rows.append(list(row))
                if len(rows) > MAX_SHEET_ROWS:
                    raise TallyParseError(f"This sheet has more than {MAX_SHEET_ROWS:,} rows.")
            if _find_header(rows) is not None:
                return _read_table(rows, "xlsx")
    finally:
        workbook.close()
    raise TallyParseError(
        "No sheet in this workbook has a Ledger (or Name) column. Export the List of Accounts from Tally."
    )


def _cell_text(value) -> str:
    """A cell as the text a person would have typed. Never a float's repr."""
    if value is None:
        return ""
    if isinstance(value, bool):
        raise _BadAmount("A yes/no cell is not an amount.")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        rounded = round(value, 2)
        if abs(value - rounded) > 1e-6:
            raise _BadAmount("The amount has more than two decimals.")
        return f"{rounded:.2f}"
    return str(value)


def _find_header(rows):
    for index, row in enumerate(rows[:30]):
        heads = [str(c).strip().casefold() if c is not None else "" for c in row]
        if any(h in _HEADERS["name"] for h in heads):
            return index, {
                key: next((i for i, h in enumerate(heads) if h in names), None)
                for key, names in _HEADERS.items()
            }
    return None


def _read_table(rows, source_format: str) -> ParsedMasters:
    found = _find_header(rows)
    if found is None:
        raise TallyParseError(
            "No Ledger (or Name) column was found. The first rows should be headings such as "
            "Ledger, Group and Opening Balance."
        )
    header_row, columns = found
    parsed = ParsedMasters(source_format=source_format)

    def cell(row, key):
        i = columns[key]
        return row[i] if i is not None and i < len(row) else None

    for offset, row in enumerate(rows[header_row + 1 :], start=header_row + 2):
        if not any(c not in (None, "") for c in row):
            continue
        if len(parsed.ledgers) + len(parsed.skipped) >= MAX_LEDGERS:
            raise TallyParseError(f"This file has more than {MAX_LEDGERS:,} ledgers.")
        raw_name = cell(row, "name")
        try:
            name = clean_text(raw_name)
            parent_text = cell(row, "group")
            parent = clean_text(parent_text, what="group name") if str(parent_text or "").strip() else ""
        except _BadText as exc:
            parsed.skipped.append(SkippedLine(offset, _shown(raw_name), str(exc)))
            continue
        try:
            opening = _table_opening(cell, row)
        except _BadAmount as exc:
            parsed.skipped.append(SkippedLine(offset, name, str(exc)))
            continue
        parsed.ledgers.append(MasterLedger(offset, name, parent, opening, None))
    return parsed


def _magnitude(text: str) -> int:
    return abs(parse_opening(text)) if text else 0


def _table_opening(cell, row) -> int:
    """Opening as debit-positive paise from either layout a list can have.

    Separate Debit and Credit columns (a bare figure is on the side of its column,
    whatever sign it carries), or one Opening column that ends in Dr or Cr or has
    a side column beside it.
    """
    debit, credit = _cell_text(cell(row, "debit")).strip(), _cell_text(cell(row, "credit")).strip()
    if debit or credit:
        owed, due = _magnitude(debit), _magnitude(credit)
        if owed and due:
            raise _BadAmount("The row has both a debit and a credit opening balance.")
        return owed - due

    amount = _cell_text(cell(row, "opening")).strip()
    if not amount:
        return 0
    side = _cell_text(cell(row, "side")).strip().casefold()
    if side and not re.search(r"(dr|cr)\s*$", amount, re.IGNORECASE):
        if side not in ("dr", "cr", "debit", "credit"):
            raise _BadAmount(f"The side {side[:10]!r} is neither Dr nor Cr.")
        amount = f"{amount} {'Dr' if side.startswith('d') else 'Cr'}"
    return parse_opening(amount)
